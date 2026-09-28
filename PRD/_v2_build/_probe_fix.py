# -*- coding: utf-8 -*-
"""修复前勘查：H4编号 / H3编号 / 目录域 / 空白标题 / 因子库正文段 / 封面与修订表"""
import docx, re, sys, zipfile
from lxml import etree
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
sys.stdout.reconfigure(encoding='utf-8')
P = r'D:\Project\TJ-nengtan\PRD\特变电工能碳数字化双中心_产品需求规格说明书_PRD_v2.0.docx'
d = docx.Document(P); z = zipfile.ZipFile(P)
kids = list(d.element.body.iterchildren())

print('=' * 76); print('【1】Heading 4 段落的编号来源（前 10 个）'); print('=' * 76)
n = 0
for p in d.paragraphs:
    if p.style.name == 'Heading 4':
        pPr = p._p.find(qn('w:pPr'))
        np_ = pPr.find(qn('w:numPr')) if pPr is not None else None
        info = '无直接numPr（依赖样式）'
        if np_ is not None:
            ni = np_.find(qn('w:numId')); lv = np_.find(qn('w:ilvl'))
            info = '直接 numPr ilvl=%s numId=%s' % (lv.get(qn('w:val')) if lv is not None else '?',
                                                    ni.get(qn('w:val')) if ni is not None else '?')
        print('   %-22s %s' % (p.text.strip()[:22], info)); n += 1
        if n >= 10: break

print('\n' + '=' * 76); print('【2】Heading 3 的 numId 分组（按 H2 归组）'); print('=' * 76)
cur2 = None
for p in d.paragraphs:
    if p.style.name == 'Heading 2': cur2 = p.text.strip()
    if p.style.name == 'Heading 3':
        pPr = p._p.find(qn('w:pPr')); np_ = pPr.find(qn('w:numPr')) if pPr is not None else None
        nid = '—'
        if np_ is not None:
            ni = np_.find(qn('w:numId'))
            nid = ni.get(qn('w:val')) if ni is not None else '?'
        print('   [%s] %-24s numId=%s' % (str(cur2)[:12], p.text.strip()[:24], nid))

print('\n' + '=' * 76); print('【3】numId 定义 & abstractNum 的 lvlText'); print('=' * 76)
num = z.read('word/numbering.xml').decode('utf-8', 'ignore')
n2a = dict(re.findall(r'<w:num w:numId="(\d+)"[^>]*>\s*<w:abstractNumId w:val="(\d+)"', num))
print('numId->abstractNumId:', {k: n2a[k] for k in sorted(n2a, key=int) if int(k) in (7, 9, 16, 27, 64, 131, 150, 161, 177)})
for aid in sorted(set(n2a.values()), key=int):
    mm = re.search(r'<w:abstractNum w:abstractNumId="%s"[^>]*>.*?</w:abstractNum>' % aid, num, re.S)
    if mm:
        lv = re.findall(r'<w:lvl w:ilvl="(\d)"[^>]*>.*?<w:lvlText w:val="([^"]*)"', mm.group(0), re.S)
        if lv and any(v for _, v in lv):
            pass
for nid in ('7', '27', '150', '177'):
    aid = n2a.get(nid)
    mm = re.search(r'<w:abstractNum w:abstractNumId="%s"[^>]*>.*?</w:abstractNum>' % aid, num, re.S) if aid else None
    if mm:
        lv0 = re.search(r'<w:lvl w:ilvl="0".*?</w:lvl>', mm.group(0), re.S)
        lt = re.search(r'<w:lvlText w:val="([^"]*)"', lv0.group(0)) if lv0 else None
        print('   numId=%-4s abstract=%s  lvl0 lvlText=%r' % (nid, aid, lt.group(1) if lt else '?'))

print('\n' + '=' * 76); print('【4】目录域 XML 片段（含 fldChar / instrText）'); print('=' * 76)
xml = z.read('word/document.xml').decode('utf-8', 'ignore')
i = xml.find('TOC')
print(xml[max(0, i - 700):i + 260].replace('><', '>\n<')[:2200])

print('\n' + '=' * 76); print('【5】空白 Heading 1 的完整 XML'); print('=' * 76)
for p in d.paragraphs:
    if p.style.name == 'Heading 1' and not p.text.strip():
        print(etree.tostring(p._p, pretty_print=True, encoding='unicode')[:1200])

print('\n' + '=' * 76); print('【6】因子库模块正文段（【n】前缀）的 pPr 样式'); print('=' * 76)
cnt = 0
for p in d.paragraphs:
    t = p.text.strip()
    if re.match(r'^【[1-8]】', t):
        cnt += 1
        if cnt <= 14:
            print('   [%s] %s' % (p.style.name, t[:44]))
print('   共 %d 个【n】开头的正文段（其中属于 Heading 4 的应被排除）' % cnt)
h4pref = sum(1 for p in d.paragraphs if p.style.name == 'Heading 4' and re.match(r'^【[1-8]】', p.text.strip()))
print('   已是 Heading 4 且带【n】前缀的段数:', h4pref)

print('\n' + '=' * 76); print('【7】封面元数据表 & 修订记录表'); print('=' * 76)
t0 = d.tables[0]
for r in t0.rows:
    print('   ', ' || '.join(c.text.strip()[:60] for c in r.cells))
print('   修订表行数:', len(d.tables[2].rows), '列数:', len(d.tables[2].columns))
last = d.tables[2].rows[-1]
print('   末行:', ' | '.join(c.text.strip()[:26] for c in last.cells))
