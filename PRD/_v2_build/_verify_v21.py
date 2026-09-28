# -*- coding: utf-8 -*-
"""v2.1 修复结果校验"""
import docx, re, sys, zipfile, collections, os
from lxml import etree
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
sys.stdout.reconfigure(encoding='utf-8')
P = r'D:\Project\TJ-nengtan\PRD\特变电工能碳数字化双中心_产品需求规格说明书_PRD_v2.1.docx'
d = docx.Document(P); z = zipfile.ZipFile(P)
xml = z.read('word/document.xml').decode('utf-8', 'ignore')
numxml = z.read('word/numbering.xml').decode('utf-8', 'ignore')

print('=' * 74); print('【A】版本与包体'); print('=' * 74)
print('文件大小 : %.2f MB' % (os.path.getsize(P) / 1048576))
print('段落 %d | 表格 %d | 图片 %d | 节 %d' % (len(d.paragraphs), len(d.tables), len(d.inline_shapes), len(d.sections)))
hs = collections.Counter(p.style.name for p in d.paragraphs if p.style.name.startswith('Heading'))
print('标题统计 :', dict(sorted(hs.items())))
print('core.version=%r | comments=%r' % (d.core_properties.version, d.core_properties.comments))

print('\n' + '=' * 74); print('【B】P0-2 空白标题'); print('=' * 74)
empties = [p.text for p in d.paragraphs if p.style.name.startswith('Heading') and not p.text.strip()]
print('空标题数 :', len(empties), '(期望 0)')
h1 = [p.text.strip() for p in d.paragraphs if p.style.name == 'Heading 1']
print('H1 数 :', len(h1), '(期望 12)')
print('H1 列表:', h1)
n7 = 0
for p in d.paragraphs:
    if p.style.name == 'Heading 1':
        pPr = p._p.find(qn('w:pPr'))
        np_ = pPr.find(qn('w:numPr')) if pPr is not None else None
        if np_ is not None and np_.find(qn('w:numId')).get(qn('w:val')) == '7':
            n7 += 1
print('共用章编号(numId=7)的 H1 数 : %d (期望 11 → 「第一章」= 顶层业务架构)' % n7)

print('\n' + '=' * 74); print('【C】P0-1 目录域'); print('=' * 74)
for m in re.finditer(r'<w:instrText[^>]*>([^<]*TOC[^<]*)</w:instrText>', xml):
    print('  指令:', repr(m.group(1)))
print('  dirty 标记数 :', xml.count('w:dirty="true"'))
print('  缓存目录条目 :', len(re.findall(r'PAGEREF', xml)))

print('\n' + '=' * 74); print('【D】P1-3 页面尺寸'); print('=' * 74)
for i, s in enumerate(d.sections):
    print('  节%d : %.2f x %.2f cm  %s' % (i, s.page_width.cm, s.page_height.cm,
          'A4 ✓' if abs(s.page_width.cm - 21.0) < .01 and abs(s.page_height.cm - 29.7) < .01 else '✗'))

print('\n' + '=' * 74); print('【E】P1-5 H3 编号与标题文本'); print('=' * 74)
miss = 0; manual = []
for p in d.paragraphs:
    if p.style.name == 'Heading 3':
        pPr = p._p.find(qn('w:pPr'))
        np_ = pPr.find(qn('w:numPr')) if pPr is not None else None
        if np_ is None:
            miss += 1; print('  ✗ 仍缺编号:', p.text.strip())
        if re.match(r'^\s*\d+\s*[\.、]', p.text.strip()):
            manual.append(p.text.strip())
print('  H3 缺编号数 : %d (期望 0)' % miss)
print('  标题残留手写序号 :', manual if manual else '无 ✓')
print('  因子库/认证/驾驶舱模块标题:')
for p in d.paragraphs:
    if p.style.name == 'Heading 3' and any(k in p.text for k in ['碳排因子', '折标煤', '认证', '驾驶舱']):
        pPr = p._p.find(qn('w:pPr')); np_ = pPr.find(qn('w:numPr'))
        print('     %-18s numId=%s' % (p.text.strip(), np_.find(qn('w:numId')).get(qn('w:val'))))

print('\n' + '=' * 74); print('【F】P1-4 因子库八段标题'); print('=' * 74)
mods = []
for p in d.paragraphs:
    if p.style.name == 'Heading 3': mods.append({'n': p.text.strip(), 'h4': []})
    elif p.style.name == 'Heading 4' and mods: mods[-1]['h4'].append(p.text.strip())
print('H4 总数 : %d (期望 312 = 280 + 32)' % sum(len(m['h4']) for m in mods))
bad = [(m['n'], len(m['h4'])) for m in mods if 4 <= len(m['h4']) < 8 and any(
    k in m['n'] for k in ['碳排因子', '折标煤', '驾驶舱', '认证', '指标'])]
for m in mods:
    if any(k in m['n'] for k in ['碳排因子', '折标煤']):
        print('     %-16s H4=%d  %s' % (m['n'], len(m['h4']), m['h4'][:3]))
gap = [m['n'] for m in mods if len(m['h4']) not in (0, 8)]
print('  非 8 段模块 :', gap if gap else '无 ✓')

print('\n' + '=' * 74); print('【G】P1-6 缺口说明 & 图片'); print('=' * 74)
notes = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith('【原型缺口】')]
print('缺口说明段 : %d (期望 3)' % len(notes))
for n in notes: print('    ', n[:64])
print('图片总数 : %d (期望 40)' % len(d.inline_shapes))

print('\n' + '=' * 74); print('【H】numbering.xml 合法性（元素顺序）'); print('=' * 74)
root = etree.fromstring(numxml.encode('utf-8'))
seq = [etree.QName(c).localname for c in root]
print('  子元素序列尾部 :', seq[-8:])
trouble = []
seen_num = False
for t in seq:
    if t == 'num': seen_num = True
    elif seen_num and t in ('abstractNum', 'numPicBullet'):
        trouble.append(t)
print('  顺序冲突(num 之后又出现 abstractNum) :', trouble if trouble else '无 ✓')
print('  numIdMacAtCleanup 位置 :',
      ('存在且位于 num 之前（非法）' if 'numIdMacAtCleanup' in seq and seq.index('numIdMacAtCleanup') < seq.index('num')
       else '正常'))
print('  num 元素数 : %d' % seq.count('num'), '| 新增 :', [n for n in re.findall(r'<w:num w:numId="(\d+)"', numxml) if int(n) >= 214])

print('\n' + '=' * 74); print('【I】修订记录表'); print('=' * 74)
for r in d.tables[2].rows[-2:]:
    print('   ', ' | '.join(c.text.strip()[:34] for c in r.cells))
print('  封面:')
for r in d.tables[0].rows:
    print('   ', ' || '.join(c.text.strip()[:46] for c in r.cells))
