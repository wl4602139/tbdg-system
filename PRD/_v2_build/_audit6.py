# -*- coding: utf-8 -*-
"""第六轮：目录编号实证 / 空单元格归属 / 修改记录 / 标题自动编号"""
import docx, re, sys, zipfile, collections
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
P = r'D:\Project\TJ-nengtan\PRD\特变电工能碳数字化双中心_产品需求规格说明书_PRD_v2.0.docx'
d = docx.Document(P); body = d.element.body
z = zipfile.ZipFile(P); xml = z.read('word/document.xml').decode('utf-8', 'ignore')

print('=' * 78); print('【1】目录中的可疑条目（重复编号 / 空条目）'); print('=' * 78)
m = re.search(r'<w:sdt>.*?</w:sdt>', xml, re.S)
seg = m.group(0) if m else ''
# 目录条目 = 每个含 PAGEREF 的段落
entries = []
for pm in re.finditer(r'<w:p\b.*?</w:p>', seg, re.S):
    s = pm.group(0)
    t = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', s))
    if 'PAGEREF' in s:
        entries.append(t)
print('目录条目数 :', len(entries))
print('\n-- 因子库/认证/驾驶舱 相关条目 --')
for e in entries:
    if any(k in e for k in ['原材料碳排因子', '电力碳排因子', '能源活动碳排因子', '折标煤', '认证资料', '认证申请', '认证结果', '碳足迹总览', '修改记录']):
        print('   >', repr(e))
print('\n-- 空/异常条目 --')
for i, e in enumerate(entries):
    if len(e.strip()) < 3:
        print('   #%d %r' % (i, e))
print('\n-- 前 8 与后 8 条 --')
for e in entries[:8]: print('   ', repr(e))
print('   ...')
for e in entries[-8:]: print('   ', repr(e))

print('\n' + '=' * 78); print('【2】标题"自动编号"定义检查'); print('=' * 78)
st = z.read('word/styles.xml').decode('utf-8', 'ignore')
for lvl in ['Heading1', 'Heading2', 'Heading3', 'Heading4']:
    mm = re.search(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % lvl, st, re.S)
    if mm:
        has = '<w:numPr>' in mm.group(0)
        print('  %-9s 绑定自动编号: %s' % (lvl, '是' if has else '否'))
num = z.read('word/numbering.xml').decode('utf-8', 'ignore') if 'word/numbering.xml' in z.namelist() else ''
print('  numbering.xml 存在:', bool(num), '| abstractNum 数:', len(re.findall(r'<w:abstractNum\b', num)))
for fm in re.finditer(r'<w:lvlText w:val="([^"]*)"', num[:12000]):
    pass
print('  多级编号样例:', re.findall(r'<w:lvlText w:val="([^"]*)"', num)[:14])

print('\n' + '=' * 78); print('【3】含空单元格的表格归属'); print('=' * 78)
cur = '(前置)'
res = []
for ch in body.iterchildren():
    if ch.tag == qn('w:p'):
        p = Paragraph(ch, d)
        if p.style.name.startswith('Heading'): cur = p.text.strip() or p.style.name
    elif ch.tag == qn('w:tbl'):
        t = Table(ch, d)
        e = sum(1 for r in t.rows for c in r.cells if not c.text.strip())
        if e:
            res.append((cur, len(t.rows), len(t.columns), e, ' | '.join(c.text.replace('\n', ' ')[:14] for c in t.rows[0].cells)))
print('含空格表数 :', len(res), '| 空格合计 :', sum(r[3] for r in res))
for c, r, co, e, h in res:
    print('   %-30s %dx%-2d 空%d  [%s]' % (c[:30], r, co, e, h[:52]))

print('\n' + '=' * 78); print('【4】修改记录表 & 目录前置'); print('=' * 78)
for i, t in enumerate(d.tables[:3]):
    print('-- 表%d  %dx%d' % (i, len(t.rows), len(t.columns)))
    for r in t.rows[:6]:
        print('     ', ' | '.join(c.text.replace('\n', ' ')[:30] for c in r.cells))

print('\n' + '=' * 78); print('【5】全文 H2 清单（41 条）'); print('=' * 78)
for i, p in enumerate([x for x in d.paragraphs if x.style.name == 'Heading 2'], 1):
    print('  %2d. %s' % (i, p.text.strip()[:80]))
