# -*- coding: utf-8 -*-
"""独立只读核验 V1.3 修订结果"""
import openpyxl, re, collections, os

P = r'D:\Project\TJ-nengtan\PRD\数据表\双中心能碳管控平台_基础参数与数据字典_V1.3.xlsx'
print('mtime:', __import__('datetime').datetime.fromtimestamp(os.path.getmtime(P)))
print('size :', os.path.getsize(P))
wb = openpyxl.load_workbook(P, data_only=True)
print('SHEETS:', wb.sheetnames)
print()

# ================= 01 基础参数表 =================
ws = wb['01-基础参数表']
rows = list(ws.iter_rows(values_only=True))
hdr = [str(x) for x in rows[0]]
print('01 表头:', hdr)

codes = collections.Counter()
for r in rows[1:]:
    c = str(r[2]).strip() if r[2] else ''
    if re.match(r'^[A-Z]+-', c):
        codes[c] += 1
dup = {k: v for k, v in codes.items() if v > 1}
print('01 参数条目:', len(codes), ' 唯一:', len(set(codes)), ' 重复编码:', dup if dup else '无')
print('EQ 组:', sorted(k for k in codes if k.startswith('EQ-')))
print('ED 组:', sorted(k for k in codes if k.startswith('ED-')))
print('ED 组条目数:', len([k for k in codes if k.startswith('ED-')]))
print()

# 02/03/04 悬挂引用反查
allcodes = set(codes)
dangle = {}

def scan(sheet, cols):
    wsx = wb[sheet]
    bad = collections.Counter()
    for i, r in enumerate(wsx.iter_rows(values_only=True), 1):
        for ci in cols:
            if ci < len(r):
                for cc in re.findall(r'\b(?:EN|OUT|VAL|CAR|EQ|ED|FAC|ORG|MD|ZP|NE|LD|GV|RV|HV|CO|OD|MT|IND)-[A-Z0-9]+\b', str(r[ci] or '')):
                    if cc.startswith('IND-'):
                        continue
                    if cc not in allcodes:
                        bad[cc] += 1
    return bad

for sh in ['02-指标清单', '03-公式变量映射', '04-工序数据项矩阵', '05-数据源清单']:
    wsx = wb[sh]
    n = wsx.max_column
    d = scan(sh, range(n))
    dangle[sh] = d
    print(f'{sh} 悬挂引用:', dict(d) if d else '无')

print()

# ================= 02 指标清单 =================
ws2 = wb['02-指标清单']
r2 = list(ws2.iter_rows(values_only=True))
txt2 = '\n'.join(' | '.join(str(v) for v in r if v is not None) for r in r2)
print('02 残留 OUT-P01:', txt2.count('OUT-P01'), ' 残留 EN-026:', txt2.count('EN-026'))
for i, r in enumerate(r2, 1):
    c = str(r[0]).strip() if r[0] else ''
    if c in ('IND-011', 'IND-012', 'IND-013', 'IND-014', 'IND-015', 'IND-016'):
        seg = str(r[7] or '')
        print(f'  row{i} {c} 依赖参数: {seg[:110]}')
print()

# ================= 04 工序矩阵 =================
ws4 = wb['04-工序数据项矩阵']
r4 = list(ws4.iter_rows(values_only=True))
print('04 max_row:', ws4.max_row, ' max_col:', ws4.max_column)
print('04 表头:', [str(x) for x in r4[0] if x is not None])
body = [r for r in r4[1:] if r[1]]
print('04 数据行:', len(body))
seq = []
for r in body:
    try:
        seq.append(int(r[0]))
    except Exception:
        pass
print('04 序号范围:', (min(seq), max(seq)) if seq else None, ' 连续:', seq == list(range(1, len(seq) + 1)))
ke_in4 = set()
km_in4 = set()
for r in body:
    for v in r[3:9]:
        for cc in re.findall(r'EN-KE\d+|OUT-KM\d+', str(v or '')):
            if cc.startswith('EN-KE'):
                ke_in4.add(cc)
            else:
                km_in4.add(cc)
    for cc in re.findall(r'EN-KE\d+|OUT-KM\d+', str(r[12] or '')):
        if cc.startswith('EN-KE'):
            ke_in4.add(cc)
        else:
            km_in4.add(cc)
ke_all = {k for k in codes if k.startswith('EN-KE')}
km_all = {k for k in codes if k.startswith('OUT-KM')}
print('EN-KE 总数:', len(ke_all), ' 04 出现:', len(ke_in4), ' 缺:', sorted(ke_all - ke_in4))
print('OUT-KM 总数:', len(km_all), ' 04 出现:', len(km_in4), ' 缺:', sorted(km_all - km_in4))
# 吨铜电耗 8 指标
tonight = ['IND-017', 'IND-019', 'IND-021', 'IND-024', 'IND-027', 'IND-029', 'IND-031', 'IND-033']
cov = set()
for r in body:
    for cc in re.findall(r'IND-\d+', str(r[12] or '')):
        cov.add(cc)
print('8 个吨铜电耗指标在04表:', {t: (t in cov) for t in tonight})
print()

# ================= 01 表 H:K 空缺 =================
miss = []
for i, r in enumerate(rows[1:], 2):
    c = str(r[2]).strip() if r[2] else ''
    if not c:
        continue
    if not r[7] or not r[8] or not r[9] or not r[10]:
        miss.append((i, c, r[7], r[8], r[9], r[10]))
print('01 表 H:K 仍有空缺的行数:', len(miss))
for m in miss[:15]:
    print('   ', m)
