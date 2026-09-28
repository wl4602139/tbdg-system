# -*- coding: utf-8 -*-
"""独立只读核验 V1.4（第 3、4、5 批修订结果）"""
import openpyxl, re, collections, os, datetime

P = r'D:\Project\TJ-nengtan\PRD\数据表\双中心能碳管控平台_基础参数与数据字典_V1.4.xlsx'
print('mtime:', datetime.datetime.fromtimestamp(os.path.getmtime(P)))
print('size :', os.path.getsize(P))
wb = openpyxl.load_workbook(P, data_only=True)
print('SHEETS(%d):' % len(wb.sheetnames), wb.sheetnames)
print()

def hdr(sheet):
    ws = wb[sheet]
    return [str(c.value) if c.value is not None else '' for c in ws[1]], ws

# ---------- 00 主数据与维度 ----------
h, ws = hdr('00-主数据与维度')
rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]
print('== 00 主数据与维度 ==')
print('  列:', h)
print('  行数:', len(rows))
codes0 = [str(r[1]).strip() for r in rows if r[1] and re.match(r'^[A-Z]+-', str(r[1]))]
print('  编码:', codes0)
print()

# ---------- 01 基础参数表 ----------
h, ws = hdr('01-基础参数表')
print('== 01 基础参数表 ==')
print('  列数:', ws.max_column, ' 列:', h)
rows = list(ws.iter_rows(min_row=2, values_only=True))
body = [r for r in rows if r[2] and re.match(r'^[A-Z]+-', str(r[2]).strip())]
print('  参数条目:', len(body))
codes = collections.Counter(str(r[2]).strip() for r in body)
dup = {k: v for k, v in codes.items() if v > 1}
print('  重复编码:', dup if dup else '无')
seq = [r[1] for r in body]
print('  序号 min/max:', min(seq), max(seq), ' 连续:', list(seq) == list(range(1, len(seq) + 1)))
# 新增列填充率
idx = {name: i for i, name in enumerate(h)}
for col in ['变量符号（原表）', '所属产线', '所属工序', '数据类型', '数据格式', '小数位',
            '单位取值规则', '取值范围', '默认值', '是否必填', '校验规则', '来源字段/表',
            '数据源编码', '关联指标编码', '参数构成/细分项']:
    if col in idx:
        i = idx[col]
        filled = sum(1 for r in body if i < len(r) and r[i] not in (None, '', '—'))
        dash = sum(1 for r in body if i < len(r) and r[i] == '—')
        print(f'    {col:<16} 实填 {filled:>3} / 破折 {dash:>3} / 空 {len(body)-filled-dash:>3}')
    else:
        print(f'    {col:<16} !! 列不存在')
# 半角等号检查
half = [(str(r[2]), str(r[i])[:40]) for r in body for i in range(len(h))
        if isinstance(r[i], str) and r[i].startswith('=')]
print('  半角等号开头文本:', len(half), half[:3] if half else '')
# 第1、2批成果保全
print('  ED 组:', sorted(k for k in codes if k.startswith('ED-')), ' EQ 组:', sorted(k for k in codes if k.startswith('EQ-')))
print()

# ---------- 02 指标清单 ----------
h, ws = hdr('02-指标清单')
r2 = list(ws.iter_rows(min_row=1, values_only=True))
print('== 02 指标清单 ==')
print('  列:', h)
inds = [r for r in r2[1:] if r[0] and re.match(r'^IND-\d+$', str(r[0]).strip())]
print('  指标数:', len(inds))
ci = h.index('规范计算公式') if '规范计算公式' in h else None
if ci is not None:
    f = [r for r in inds if r[ci] and str(r[ci]).strip()]
    fullw = [r for r in f if str(r[ci]).startswith('＝')]
    halfw = [r for r in f if str(r[ci]).startswith('=')]
    print('  公式非空:', len(f), ' 全角开头:', len(fullw), ' 半角开头:', len(halfw))
    for r in f[:3]:
        print('     ', r[0], '→', str(r[ci])[:70])
else:
    print('  !! 找不到「规范计算公式」列')
grp = [r for r in r2[1:] if r[0] and not re.match(r'^IND-\d+$', str(r[0]).strip()) and r[0]]
print('  A 列非指标行:', [(i, str(r[0])) for i, r in enumerate(r2[1:], 2) if r[0] and not re.match(r'^IND-', str(r[0]))] or '无')
print('  分组列取值:', collections.Counter(str(r[2]) for r in r2[1:] if len(r) > 2 and r[2]).most_common(8))
print()

# ---------- 03 ----------
h, ws = hdr('03-公式变量映射')
print('== 03 公式变量映射 ==')
print('  列:', h, ' 行数:', ws.max_row - 1)
gi = h.index('颗粒度层级') if '颗粒度层级' in h else None
if gi is not None:
    print('  颗粒度层级分布:', dict(collections.Counter(str(r[gi]) for r in ws.iter_rows(min_row=2, values_only=True) if r[gi])))
print()

# ---------- 04 ----------
h, ws = hdr('04-工序数据项矩阵')
b4 = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
print('== 04 工序数据项矩阵 ==')
print('  max_column:', ws.max_column, ' 数据行:', len(b4))
print('  列:', h)
ke = set()
for r in b4:
    for cc in re.findall(r'EN-KE\d+', ' '.join(str(x) for x in r)): ke.add(cc)
ke_all = {k for k in codes if k.startswith('EN-KE')}
print('  EN-KE 覆盖:', len(ke), '/', len(ke_all))
ton = ['IND-017', 'IND-019', 'IND-021', 'IND-024', 'IND-027', 'IND-029', 'IND-031', 'IND-033']
cov = set()
for r in b4:
    ji = h.index('关联指标') if '关联指标' in h else 13
    for cc in re.findall(r'IND-\d+', str(r[ji] or '')): cov.add(cc)
print('  8 个吨铜电耗指标:', {t: (t in cov) for t in ton})
print()

# ---------- 05 ----------
h, ws = hdr('05-数据源清单')
print('== 05 数据源清单 ==')
print('  列:', h)
src = [str(r[1]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
print('  SRC:', src)
nums = sorted(int(re.sub(r'\D', '', s)) for s in src if re.match(r'^SRC-\d+$', s))
print('  SRC 连续:', nums == list(range(1, len(nums) + 1)), nums)
print()

# ---------- 08 ----------
h, ws = hdr('08-枚举字典')
rows8 = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]
print('== 08 枚举字典 ==')
print('  列:', h, ' 条目:', len(rows8), ' max_row:', ws.max_row)
blank = [i for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2) if not any(r)]
print('  空洞行:', blank[:10] if blank else '无')
print('  枚举组:', collections.Counter(str(r[0]) for r in rows8).most_common(30))
print()

# ---------- 09 ----------
h, ws = hdr('09-指标参数构成明细')
rows9 = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]
print('== 09 指标参数构成明细 ==')
print('  列:', h, ' 条目:', len(rows9), ' max_row:', ws.max_row)
blank9 = [i for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2) if not any(r)]
print('  空洞行:', blank9[:10] if blank9 else '无')
inds9 = set()
for r in rows9:
    for cc in re.findall(r'IND-\d+', ' '.join(str(x) for x in r)): inds9.add(cc)
print('  覆盖指标数:', len(inds9), '/ 93')
print()

# ---------- 悬挂引用（全表） ----------
print('== 悬挂引用反查 ==')
allcodes = set(codes) | set(codes0)
CODE = re.compile(r'\b(?:EN|OUT|VAL|CAR|EQ|ED|FAC|ORG|MD|ZP|NE|LD|GV|RV|HV|CO|OD|MT|SRC)-[A-Z0-9]+\b')
for sh in wb.sheetnames:
    wsx = wb[sh]
    bad = collections.Counter()
    for r in wsx.iter_rows(values_only=True):
        for v in r:
            if isinstance(v, str):
                for cc in CODE.findall(v):
                    if cc.startswith('SRC-'):
                        continue
                    if cc not in allcodes:
                        bad[cc] += 1
    print(f'  {sh}: {dict(bad) if bad else "无"}')
