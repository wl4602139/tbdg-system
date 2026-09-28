# -*- coding: utf-8 -*-
"""生成《双中心能碳管控平台_基础参数与指标字典_V1.0.xlsx》"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import xlrd
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "xlrd>=2.0"])
    import xlrd
try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from data_params import PARAMS
from data_dicts import SOURCES, ENUMS

BASE_DIR = r"D:\Project\TJ-nengtan\PRD\数据表"
SRC = os.path.join(BASE_DIR, "“双中心”项目能碳管控指标体系（9.16）V1.7.et")
OUT = os.path.join(BASE_DIR, "双中心能碳管控平台_基础参数与指标字典_V1.0.xlsx")
TODAY = "2026-09-18"


def xl_color(css_hex):
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError("Expected #RRGGBB, got: %s" % css_hex)
    return "FF" + value


XL_HEAD_BG = xl_color("#4472C4")
XL_HEAD_FG = xl_color("#FFFFFF")
XL_TITLE_BG = xl_color("#2F5597")
XL_LIGHT = xl_color("#D9E2F3")
XL_INPUT = xl_color("#FFF2CC")
XL_CALC = xl_color("#EDEDED")
XL_BORDER = xl_color("#BFBFBF")
XL_ALT = xl_color("#DEEAF6")
XL_SECT = xl_color("#EAF0FA")

thin_side = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
F_TITLE = Font(name="微软雅黑", size=14, bold=True, color=XL_HEAD_FG)
F_SUB = Font(name="微软雅黑", size=10, color="FF404040")
F_HEAD = Font(name="微软雅黑", size=10, bold=True, color=XL_HEAD_FG)
F_BODY = Font(name="微软雅黑", size=10)
F_BOLD = Font(name="微软雅黑", size=10, bold=True)
F_SECT = Font(name="微软雅黑", size=11, bold=True, color=XL_TITLE_BG)
F_MONO = Font(name="Consolas", size=10)
A_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_LEFT_C = Alignment(horizontal="left", vertical="center", wrap_text=True)
FILL_HEAD = PatternFill("solid", fgColor=XL_HEAD_BG)
FILL_TITLE = PatternFill("solid", fgColor=XL_TITLE_BG)
FILL_SECT = PatternFill("solid", fgColor=XL_SECT)
FILL_INPUT = PatternFill("solid", fgColor=XL_INPUT)
FILL_CALC = PatternFill("solid", fgColor=XL_CALC)
FILL_LIGHT = PatternFill("solid", fgColor=XL_LIGHT)

# ============================================================
# 1. 读取源指标体系
# ============================================================
sht = xlrd.open_workbook(SRC).sheet_by_index(0)


def cv(r, c):
    v = sht.cell_value(r, c)
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v).strip()


def clean(text):
    return re.sub(r"[ \t]+", " ", text.replace("\r", "")).strip()


INDICATORS = []
SECTION = ""
for r in range(2, sht.nrows):
    cols = [cv(r, c) for c in range(sht.ncols)]
    if cols[1] and not cols[2]:
        SECTION = cols[1]
        continue
    if not cols[1] or not cols[2]:
        continue
    INDICATORS.append({
        "seq": int(float(cols[1])), "name": clean(cols[2]), "section": SECTION,
        "definition": clean(cols[3]), "note": cols[5].replace("\r", ""),
        "unit": clean(cols[6]), "center": clean(cols[7]), "cat": clean(cols[8]),
        "src": clean(cols[9]), "usage": clean(cols[10]),
        "scope": cols[11].replace("\r", ""), "need": cols[12].replace("\r", ""),
        "work": cols[13].replace("\r", ""),
        "remark": cols[14].replace("\r", "") if len(cols) > 14 else "",
    })
assert len(INDICATORS) == 94

# ---- 公式变量解析 ----
VAR_MAP_EXACT = {
    "E": "EN-013", "n": "EN-014", "Ei": "EN-015", "ki": "EN-016",
    "C": "CAR-001", "C燃烧": "CAR-002", "C过程": "CAR-003", "C购入电": "CAR-004",
    "C输出电": "CAR-005", "C购入热": "CAR-006", "C输出热": "CAR-007", "C回收利用": "CAR-008",
    "I": "IND-005", "r": "IND-006", "R": "EN-017", "Eui": "IND-007", "Ez": "EN-005",
    "Q": "EN-020", "Enva": "IND-008", "Gnva": "VAL-001", "S": "IND-009",
    "Res": "EQ-001", "Ets": "EQ-002", "g": "IND-010", "G": "VAL-002",
    "Rcf": "IND-002", "Ncf": "CAR-010", "N": "CAR-011",
    "q电": "IND-013", "Q电": "EN-001", "q蒸汽": "IND-014", "Q蒸汽": "EN-007",
    "q天然气": "IND-015", "Q天然气": "EN-006", "q水": "IND-016", "Q水": "EN-026",
    "M": "OUT-008", "e": "IND-012",
}
FORMULA_OVERRIDE = {
    1: "CF = Σ(活动数据i × 排放因子i)（摇篮到大门阶段）",
    3: "E = Σ(Ei × ki)   (i = 1..n)",
    4: "C = C燃烧 + C过程 + C购入电 + C购入热 − C输出电 − C输出热 − C回收利用",
    11: "W = Σ各计量点用水量（生产用水 + 生活用水）",
}


def parse_vars(note):
    if not note:
        return []
    out = []
    for seg in re.split(r"[；;\n]", note.replace(" / ", "；")):
        seg = seg.strip().rstrip("；;").strip()
        m = re.match(r"^([^：:]{1,45})[：:](.*)$", seg, re.S)
        if not m:
            continue
        code = m.group(1).strip().replace("=", "-").replace(" ", "")
        if "。" in code or not re.match(r"^[A-Za-z]", code):
            continue
        desc = clean(m.group(2)).rstrip("；;").strip()
        um = re.search(r"[，,；;（(]单位[为是]?\s*([^，。；)]+)", desc)
        if not um:
            um = re.search(r"^单位[为是]?\s*([^，。；)]+)", desc)
        unit = um.group(1).strip() if um else ""
        if unit and not re.search(r"[A-Za-z%]", unit) and len(unit) > 4:
            unit = ""
        meaning = re.sub(r"[，,；;（(]单位[为是]?[^，。；)]+[，。；)）]?", "", desc).strip("，。 ")
        out.append({"sym": code, "desc": desc, "meaning": meaning, "unit": unit})
    return out


def map_param(ind, var, idx):
    if idx == 0:
        return "IND-%03d" % ind["seq"]
    sym = var["sym"]
    if sym in VAR_MAP_EXACT:
        return VAR_MAP_EXACT[sym]
    if sym[0] == "Q":
        for k, code in (("蒸汽", "EN-022"), ("天然气", "EN-023"), ("热力", "EN-024"), ("电", "EN-021")):
            if k in sym:
                return code
        return "EN-025"
    if sym[0] == "E":
        return "EN-025"
    if sym[0] in ("q", "u", "e"):
        return "IND-%03d" % ind["seq"]
    if sym[0] == "M":
        return "OUT-008"
    if sym[0] == "G":
        return "VAL-012"
    return ""


def build_formula(ind, vs):
    if ind["seq"] in FORMULA_OVERRIDE:
        return FORMULA_OVERRIDE[ind["seq"]]
    if len(vs) == 3:
        return "%s = %s / %s" % (vs[0]["sym"], vs[1]["sym"], vs[2]["sym"])
    if len(vs) == 2:
        return "%s = %s" % (vs[0]["sym"], vs[1]["sym"])
    return "见指标说明"


def clean_name(name):
    parts = name.split("，")
    if len(parts) > 1 and ("建议" in parts[-1] or "去掉" in parts[-1]):
        return parts[0]
    return name


def fmt_of(unit):
    u = unit or ""
    if u == "%":
        return "百分比", "2"
    if "tCO2/台套" in u or "tCO2/km" in u:
        return "数值（4 位小数）", "4"
    if u.startswith("tCO2") or u == "tce" or u.startswith("tce"):
        return "数值（3 位小数）", "3"
    if u == "kWh":
        return "数值（0 位小数）", "0"
    return "数值（2 位小数）", "2"


ROLE_MAP = {"ORG": "维度", "MD": "维度", "EN": "分子·能源消费量", "OUT": "分母·产量",
            "VAL": "分母·产值", "EQ": "分子·装备额定功率", "CAR": "加项·碳排放分量",
            "FAC": "计算系数", "IND": "指标值（被计算量）"}

PARAM_NAME = {p[0]: p[1] for p in PARAMS}
PARAM_CAT = {p[0]: p[3] for p in PARAMS}
PARAM_UNIT = {p[0]: p[5] for p in PARAMS}


def role_of(ind, var, param, idx, nvars):
    if idx == 0:
        return "指标值（被计算量）"
    pname = PARAM_NAME.get(param, "") or "—"
    pcat = PARAM_CAT.get(param, "")
    if ind["ind_code"] == "IND-003":
        return {1: "维度（能源品种计数）", 2: "因子（分品种能源消费量）",
                3: "系数（分品种折标准煤系数）"}.get(idx, "参数（%s）" % pname)
    if ind["ind_code"] == "IND-004":
        if any(k in var["sym"] for k in ("输出电", "输出热", "回收利用")):
            return "扣减项（%s）" % pname
        return "加项（%s）" % pname
    if nvars == 3:
        return ("分子 ｜ " if idx == 1 else "分母 ｜ ") + pname
    if pcat == "计算系数":
        return "系数（%s）" % pname
    if pcat in ("组织主数据", "生产主数据", "统计维度"):
        return "维度（%s）" % pname
    return "参数（%s）" % pname


for ind in INDICATORS:
    ind["ind_code"] = "IND-%03d" % ind["seq"]
    ind["vars"] = parse_vars(ind["note"])
    ind["formula"] = build_formula(ind, ind["vars"])
    ind["short_name"] = clean_name(ind["name"])
    ind["name_note"] = ""
    if ind["short_name"] != ind["name"]:
        ind["name_note"] = "原表指标名称含说明「%s」。" % ind["name"][len(ind["short_name"]):].lstrip("，")
    for i, v in enumerate(ind["vars"]):
        v["param"] = map_param(ind, v, i)
    codes = []
    for v in ind["vars"]:
        if v["param"] and not v["param"].startswith("IND-") and v["param"] not in codes:
            codes.append(v["param"])
    ind["param_codes"] = codes

# ============================================================
# 2. 组装「基础参数表」统一表（参数 + 指标）
# ============================================================
UNIFORM = []
for p in PARAMS:
    UNIFORM.append([p[0], p[1], "参数"] + p[2:])


def ind_to_uniform(ind):
    fmt, dec = fmt_of(ind["unit"])
    center_lbl = "集采中心" if ind["center"] == "集采" else "集控中心"
    syms = "、".join(v["sym"] for v in ind["vars"]) or "—"
    freq = "月度（年度汇算）" if ("年度" in ind["definition"] or "年度" in ind["note"]) else "月度"
    desc = (ind["definition"] or "—").rstrip("。； ")
    if ind["formula"] != "见指标说明":
        desc = desc + "；计算公式：" + ind["formula"]
    remark_parts = [ind["name_note"]]
    if ind["src"]:
        remark_parts.append("指标来源：%s。" % ind["src"])
    if ind["remark"]:
        remark_parts.append("原表备注：%s" % ind["remark"])
    remark_parts.append("覆盖范围、指标用途及各单位实施要求详见「03-指标清单」。")
    return [ind["ind_code"], ind["short_name"], "指标", syms, "指标 / " + (ind["cat"] or "—"),
            desc, ind["unit"], fmt, dec, "≥0", "—", "%s（系统自动计算）" % center_lbl,
            "公式变量：" + syms, "自动计算", freq, "是", "电装·%s" % center_lbl,
            "、".join(ind["param_codes"]) or "—", "≥0；分母为空或 0 时指标置空",
            "".join(remark_parts)]


for ind in INDICATORS:
    UNIFORM.append(ind_to_uniform(ind))

UNIFORM.sort(key=lambda x: x[0])

# ============================================================
# 3. 建立工作簿
# ============================================================
wb = openpyxl.Workbook()
wb.properties.title = "双中心能碳管控平台 基础参数与指标字典 V1.0"
BASE_FONT = "微软雅黑"


def style_header(ws, row, ncol, start=1):
    for c in range(start, start + ncol):
        cell = ws.cell(row=row, column=c)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = A_CENTER
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fill_table(ws, headers, rows, widths, wrap_cols=(), freeze="D2", filt=True, headrow=2):
    for i, h in enumerate(headers, start=1):
        ws.cell(row=headrow, column=i, value=h)
    style_header(ws, headrow, len(headers))
    for ri, row in enumerate(rows, start=headrow + 1):
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = F_BODY
            cell.border = BORDER
            cell.alignment = A_LEFT if ci in wrap_cols else A_LEFT_C
            if isinstance(val, float):
                cell.number_format = "#,##0.00"
    set_widths(ws, widths)
    if freeze:
        ws.freeze_panes = freeze
    if filt and rows:
        ws.auto_filter.ref = "A%d:%s%d" % (headrow, get_column_letter(len(headers)), headrow + len(rows))


# ---------- Sheet 03-指标清单 ----------
ws3 = wb.active
ws3.title = "03-指标清单"
H3 = ["序号", "指标编码", "指标名称", "指标层级", "指标定义", "计算公式", "指标说明（公式变量定义）",
      "指标单位", "所属中心", "指标类别", "指标来源", "指标用途", "覆盖范围", "所需各单位数据",
      "涉及参数编码", "各单位所需开展工作", "备注"]
W3 = [6, 11, 34, 26, 44, 40, 58, 18, 9, 10, 20, 44, 40, 40, 34, 50, 26]
rows3 = []
for ind in INDICATORS:
    remark = ind["remark"]
    if ind["name_note"]:
        remark = (ind["name_note"] + remark).strip()
    rows3.append([ind["seq"], ind["ind_code"], ind["name"], ind["section"], ind["definition"],
                  ind["formula"], ind["note"].replace("\n", " "), ind["unit"], ind["center"],
                  ind["cat"], ind["src"], ind["usage"], ind["scope"].replace("\n", " / "),
                  ind["need"].replace("\n", " / "), "、".join(ind["param_codes"]) or "—",
                  ind["work"].replace("\n", " / "), remark])
ws3["A1"] = "03-指标清单 · 共 94 项指标（源自《“双中心”项目能碳管控指标体系（9.16）V1.7》）"
ws3["A1"].font = F_TITLE
ws3["A1"].fill = FILL_TITLE
ws3["A1"].alignment = A_LEFT_C
ws3.merge_cells("A1:Q1")
ws3.row_dimensions[1].height = 24
fill_table(ws3, H3, rows3, W3, wrap_cols=(3, 4, 5, 6, 7, 11, 12, 13, 14, 15, 16, 17), freeze="D3", headrow=2)
ws3.auto_filter.ref = "A2:Q%d" % (2 + len(rows3))

# ---------- Sheet 02-基础参数表 ----------
ws2 = wb.create_sheet("02-基础参数表")
H2 = ["序号", "编码", "名称", "类型", "符号", "参数类别", "数据说明", "单位", "数据格式", "小数位",
      "取值范围 / 枚举", "默认取值 / 参考值", "数据来源", "来源对象 / 字段", "获取方式", "更新频率",
      "必填", "责任单位", "关联编码", "校验规则", "备注"]
W2 = [6, 12, 26, 8, 16, 14, 62, 18, 15, 8, 26, 18, 28, 28, 18, 18, 7, 20, 34, 24, 62]
rows2 = [[i] + row for i, row in enumerate(UNIFORM, start=1)]
ws2["A1"] = "02-基础参数表 · 参数 %d 项 + 指标 %d 项 = %d 项（可用筛选按「类型 / 参数类别」查看）" % (
    len(PARAMS), len(INDICATORS), len(UNIFORM))
ws2["A1"].font = F_TITLE
ws2["A1"].fill = FILL_TITLE
ws2["A1"].alignment = A_LEFT_C
ws2.merge_cells("A1:U1")
ws2.row_dimensions[1].height = 24
fill_table(ws2, H2, rows2, W2, wrap_cols=(3, 7, 11, 13, 14, 19, 21), freeze="D3", headrow=2)
ws2.auto_filter.ref = "A2:U%d" % (2 + len(rows2))
last2 = 2 + len(rows2)
ws2.conditional_formatting.add(
    "A3:U%d" % last2,
    FormulaRule(formula=['$D3="指标"'], fill=PatternFill("solid", fgColor=XL_ALT), stopIfTrue=False))
dv = DataValidation(type="list", formula1="='02-基础参数表'!$B$3:$B$%d" % last2, allow_blank=True)
ws2.add_data_validation(dv)

# ---------- Sheet 04-公式变量映射 ----------
ws4 = wb.create_sheet("04-公式变量映射")
H4 = ["指标编码", "指标名称", "变量序号", "变量符号", "变量含义", "变量角色", "对应参数编码", "变量单位", "原表变量定义"]
W4 = [11, 34, 9, 24, 42, 20, 14, 22, 60]
rows4 = []
for ind in INDICATORS:
    if not ind["vars"]:
        rows4.append([ind["ind_code"], ind["short_name"], 0, "—", "无显式变量，按外部标准方法核算",
                      "指标值（被计算量）", ind["ind_code"], ind["unit"], ind["note"].replace("\n", " ")])
        continue
    for i, v in enumerate(ind["vars"], start=1):
        vunit = v["unit"] or PARAM_UNIT.get(v["param"], "") or (ind["unit"] if v["param"].startswith("IND-") else "") or "—"
        rows4.append([ind["ind_code"], ind["short_name"], i, v["sym"], v["meaning"] or v["desc"],
                      role_of(ind, v, v["param"], i - 1, len(ind["vars"])), v["param"] or "—",
                      vunit, v["desc"]])
ws4["A1"] = "04-公式变量映射 · 指标计算公式中的每个变量 ↔ 基础参数表条目（用于公式驱动的自动取数与说明查询）"
ws4["A1"].font = F_TITLE
ws4["A1"].fill = FILL_TITLE
ws4["A1"].alignment = A_LEFT_C
ws4.merge_cells("A1:I1")
ws4.row_dimensions[1].height = 24
fill_table(ws4, H4, rows4, W4, wrap_cols=(2, 5, 9), freeze="C3", headrow=2)

# ---------- Sheet 05-数据源清单 ----------
ws5 = wb.create_sheet("05-数据源清单")
H5 = ["来源编码", "数据源系统", "归属方", "提供数据内容", "对接方式", "更新频率", "责任方", "备注"]
W5 = [10, 26, 16, 62, 20, 18, 20, 46]
ws5["A1"] = "05-数据源清单 · 参数的取数系统、对接方式与责任分工"
ws5["A1"].font = F_TITLE
ws5["A1"].fill = FILL_TITLE
ws5["A1"].alignment = A_LEFT_C
ws5.merge_cells("A1:H1")
ws5.row_dimensions[1].height = 24
fill_table(ws5, H5, SOURCES, W5, wrap_cols=(2, 4, 8), freeze="A3", headrow=2)

# ---------- Sheet 06-枚举字典 ----------
ws6 = wb.create_sheet("06-枚举字典")
H6 = ["字典编码", "字典名称", "枚举值", "字典说明", "适用字段", "备注"]
W6 = [10, 20, 46, 40, 42, 42]
ws6["A1"] = "06-枚举字典 · 全部枚举字段的取值域"
ws6["A1"].font = F_TITLE
ws6["A1"].fill = FILL_TITLE
ws6["A1"].alignment = A_LEFT_C
ws6.merge_cells("A1:F1")
ws6.row_dimensions[1].height = 24
fill_table(ws6, H6, ENUMS, W6, wrap_cols=(3, 4, 5, 6), freeze="A3", headrow=2)

# ---------- Sheet 01-说明与查询 ----------
ws1 = wb.create_sheet("01-说明与查询", 0)
set_widths(ws1, [26, 30, 22, 26, 18, 22, 18, 18])
ws1["A1"] = "双中心能碳管控平台 · 基础参数与指标字典"
ws1["A1"].font = Font(name=BASE_FONT, size=16, bold=True, color=XL_HEAD_FG)
ws1["A1"].fill = FILL_TITLE
ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws1.merge_cells("A1:H1")
ws1.row_dimensions[1].height = 32
ws1["A2"] = "依据《“双中心”项目能碳管控指标体系（9.16）V1.7》提炼 · 用于交付开发阶段的参数与数据说明速查"
ws1["A2"].font = F_SUB
ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws1.merge_cells("A2:H2")
ws1["A3"] = "版本 V1.0　|　生成日期 %s　|　指标 %d 项　|　原子参数 %d 项　|　变量映射 %d 行" % (
    TODAY, len(INDICATORS), len(PARAMS), len(rows4))
ws1["A3"].font = F_SUB
ws1["A3"].alignment = Alignment(horizontal="center", vertical="center")
ws1.merge_cells("A3:H3")

# 查询区
ws1["A5"] = "① 快速查询　—　在 B6 输入或下拉选择「编码 / 名称」，下方自动返回该参数或指标的全部说明信息"
ws1["A5"].font = F_SECT
ws1.merge_cells("A5:H5")
ws1["A6"] = "查询键（编码或名称）"
ws1["A6"].font = F_BOLD
ws1["B6"] = "EN-013"
ws1["B6"].fill = FILL_INPUT
ws1["B6"].font = F_BOLD
ws1["B6"].border = BORDER
ws1["C6"] = "查询状态"
ws1["C6"].font = F_BOLD
ws1["D6"] = '=IF($B$6="","请输入查询键（编码或名称）",IF($B$7="","未找到匹配记录，请检查编码或名称","匹配成功"))'
ws1["D6"].font = F_BODY
ws1["D6"].border = BORDER
ws1.merge_cells("D6:H6")
ws1["A7"] = "匹配行号（辅助）"
ws1["A7"].font = F_BODY
ws1["B7"] = "=IF($B$6=\"\",\"\",IFERROR(MATCH($B$6,'02-基础参数表'!$B:$B,0),IFERROR(MATCH($B$6,'02-基础参数表'!$C:$C,0),\"\")))"
ws1["B7"].font = F_BODY
ws1["B7"].fill = FILL_CALC
ws1["B7"].border = BORDER
ws1["D7"] = "在 WPS / Excel 中修改 B6 后会自动重算；单元格内已预置示例（EN-013 综合能源消费量）的缓存结果。"
ws1["D7"].font = F_SUB
ws1["D7"].alignment = A_LEFT_C
ws1.merge_cells("D7:H7")

QUERY_FIELDS = [
    ("编码", "B"), ("名称", "C"), ("类型（参数 / 指标）", "D"), ("符号", "E"), ("参数类别", "F"),
    ("数据说明", "G"), ("单位", "H"), ("数据格式", "I"), ("小数位", "J"),
    ("取值范围 / 枚举", "K"), ("默认取值 / 参考值", "L"), ("数据来源", "M"),
    ("来源对象 / 字段", "N"), ("获取方式", "O"), ("更新频率", "P"), ("是否必填", "Q"),
    ("责任单位", "R"), ("关联编码", "S"), ("校验规则", "T"), ("备注", "U"),
]
COLMAP = {ch: i for i, ch in enumerate("ABCDEFGHIJKLMNOPQRSTU", start=1)}
start = 8
for i, (label, col) in enumerate(QUERY_FIELDS):
    r = start + i
    ws1.cell(row=r, column=1, value=label).font = F_BOLD
    ws1.cell(row=r, column=1).border = BORDER
    ws1.cell(row=r, column=1).fill = FILL_LIGHT
    ws1.cell(row=r, column=1).alignment = A_LEFT_C
    c = ws1.cell(row=r, column=2,
                 value='=IF(OR($B$6="",$B$7=""),"",INDEX(\'02-基础参数表\'!$A:$U,$B$7,%d))' % COLMAP[col])
    c.font = F_BODY
    c.border = BORDER
    c.alignment = A_LEFT
    ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    if label in ("数据说明", "取值范围 / 枚举", "来源对象 / 字段", "校验规则", "备注"):
        ws1.row_dimensions[r].height = 30

dv_q = DataValidation(type="list", formula1="='02-基础参数表'!$B$3:$B$%d" % last2, allow_blank=True)
ws1.add_data_validation(dv_q)
dv_q.add(ws1["B6"])

# 工作表说明
r0 = start + len(QUERY_FIELDS) + 2
ws1.cell(row=r0, column=1, value="② 工作表说明").font = F_SECT
ws1.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=8)
head_sheet = ["工作表", "内容", "主要用途", "条目数"]
for i, h in enumerate(head_sheet, start=1):
    ws1.cell(row=r0 + 1, column=i, value=h)
style_header(ws1, r0 + 1, 4)
sheet_desc = [
    ["01-说明与查询", "工作簿说明、快速查询区、编码规则、常用查询公式示例", "入口页：速查任意参数 / 指标的说明、来源、格式", "1"],
    ["02-基础参数表", "参数与指标的统一定义表（含来源、格式、校验、责任单位）", "公式驱动速查的主表，支持按类型 / 类别筛选", "%d 项" % len(UNIFORM)],
    ["03-指标清单", "94 项指标的原始定义、计算公式、层级、类别、来源、用途、覆盖范围、实施要求", "保留指标体系原文信息，供需求与验收对照", "%d 项" % len(INDICATORS)],
    ["04-公式变量映射", "指标公式中的每个变量 ↔ 基础参数表条目", "支撑公式自动取数、口径核对与影响分析", "%d 行" % len(rows4)],
    ["05-数据源清单", "各数据源系统、归属方、提供内容、对接方式、责任方", "接口与取数方案的设计依据", "%d 项" % len(SOURCES)],
    ["06-枚举字典", "组织、产线、工序、能源品种、格式、频率等全部枚举值域", "字段取值约束与主数据台账", "%d 项" % len(ENUMS)],
]
for i, row in enumerate(sheet_desc):
    r = r0 + 2 + i
    for j, val in enumerate(row, start=1):
        cell = ws1.cell(row=r, column=j, value=val)
        cell.font = F_BODY
        cell.border = BORDER
        cell.alignment = A_LEFT
ws1.merge_cells(start_row=r0 + 2, start_column=2, end_row=r0 + 2, end_column=3)

# 编码规则
r1 = r0 + 2 + len(sheet_desc) + 2
ws1.cell(row=r1, column=1, value="③ 编码规则（「02-基础参数表」编码列前缀）").font = F_SECT
ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=8)
codes = [("ORG-", "组织主数据", "经营单位、项目公司、生产管理单元、计量点、绑定关系"),
         ("MD-", "生产主数据", "产品型号、产品种类、产线、工序、映射关系、生产工单、统计期"),
         ("EN-", "能源数据", "电 / 天然气 / 蒸汽 / 热力 / 油品 / 水等消耗量，综合能源消费量，折标系数引用"),
         ("OUT-", "产量数据", "产品产量、工序产量、订单产量（变压器 / 线缆 / 铁心 / 电容器等）"),
         ("VAL-", "产值数据", "工业增加值、产品产值、各工序对应产值"),
         ("CAR-", "碳排放数据", "二氧化碳排放量及其各项分量、产品碳足迹、产品类别数量"),
         ("EQ-", "设备数据", "节能装备与全部装备的额定功率及台账"),
         ("FAC-", "计算系数", "折标准煤系数、碳排放因子、含碳量、低位发热量、碳氧化率"),
         ("IND-", "指标", "94 项能碳管控指标本体（IND-001 ~ IND-094）"),
         ("SRC-", "数据源", "数据源系统（见「05-数据源清单」）")]
for i, (pre, name, desc) in enumerate(codes):
    r = r1 + 1 + i
    ws1.cell(row=r, column=1, value=pre).font = F_BOLD
    ws1.cell(row=r, column=1).border = BORDER
    ws1.cell(row=r, column=2, value=name).font = F_BODY
    ws1.cell(row=r, column=2).border = BORDER
    ws1.cell(row=r, column=3, value=desc).font = F_BODY
    ws1.cell(row=r, column=3).border = BORDER
    ws1.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)

# 示例公式
r2 = r1 + 1 + len(codes) + 2
ws1.cell(row=r2, column=1, value="④ 常用查询公式示例（复制后在最前补一个等号即可使用）").font = F_SECT
ws1.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=8)
samples = [
    ("按编码查数据说明", 'VLOOKUP("EN-013",\'02-基础参数表\'!$B:$U,6,0)'),
    ("按编码查数据来源", 'VLOOKUP("EN-013",\'02-基础参数表\'!$B:$U,12,0)'),
    ("按编码查数据格式", 'VLOOKUP("IND-035",\'02-基础参数表\'!$B:$U,8,0)'),
    ("按编码查单位与校验", 'VLOOKUP("VAL-001",\'02-基础参数表\'!$B:$U,7,0)&" / "&VLOOKUP("VAL-001",\'02-基础参数表\'!$B:$U,19,0)'),
    ("按名称查数据来源", 'VLOOKUP("综合能源消费量",\'02-基础参数表\'!$C:$U,11,0)'),
    ("编码与名称双向兼容", 'IFNA(VLOOKUP($B$6,\'02-基础参数表\'!$B:$U,6,0),VLOOKUP($B$6,\'02-基础参数表\'!$C:$U,5,0))'),
    ("查指标计算公式", 'VLOOKUP("IND-004",\'03-指标清单\'!$B:$Q,5,0)'),
    ("查指标覆盖范围", 'VLOOKUP("IND-017",\'03-指标清单\'!$B:$Q,12,0)'),
    ("查某指标用到的参数编码", 'VLOOKUP("IND-035",\'03-指标清单\'!$B:$Q,14,0)'),
    ("查某指标变量的参数归属", 'VLOOKUP("IND-035",\'04-公式变量映射\'!$A:$I,7,0)'),
    ("统计某类参数数量", 'COUNTIF(\'02-基础参数表\'!$F:$F,"能源数据")'),
    ("按参数类别列出全部名称", 'TEXTJOIN("、",1,IF(\'02-基础参数表\'!$F$3:$F$%d="能源数据",\'02-基础参数表\'!$C$3:$C$%d,""))' % (last2, last2)),
]
for i, (scene, fx) in enumerate(samples):
    r = r2 + 1 + i
    ws1.cell(row=r, column=1, value=scene).font = F_BODY
    ws1.cell(row=r, column=1).border = BORDER
    ws1.cell(row=r, column=2, value=fx).font = F_MONO
    ws1.cell(row=r, column=2).border = BORDER
    ws1.cell(row=r, column=2).alignment = A_LEFT_C
    ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)

# 维护约定
r3 = r2 + 1 + len(samples) + 2
ws1.cell(row=r3, column=1, value="⑤ 维护约定").font = F_SECT
ws1.merge_cells(start_row=r3, start_column=1, end_row=r3, end_column=8)
notes = [
    "1. 本表为唯一参数口径来源：指标计算公式中的变量一律通过「04-公式变量映射」关联到「02-基础参数表」条目，不得另立口径。",
    "2. 折标准煤系数、碳排放因子（FAC-*）由电装统一维护，年度复核；参考值仅供开发校验，正式取值以电装维护值为准。",
    "3. 分母类参数（产量、产值、装备功率）在公式中必须做空值与 0 值保护，分母为空或 0 时指标返回空值。",
    "4. 集控中心侧的能耗类数据统一由系统计算；手工录入仅覆盖绿电 / 绿证、难以自动采集的油品与天然气、工业增加值、用水量、装备台账。",
    "5. 新增指标或参数时，须同步补充：编码与名称、数据说明、单位、数据格式、数据来源、获取方式、更新频率、责任单位，并在「04-公式变量映射」中登记变量归属。",
    "6. 指标定义、覆盖范围、各单位实施要求以「03-指标清单」原文为准；本表不改变原指标的核算口径。",
]
for i, t in enumerate(notes):
    r = r3 + 1 + i
    ws1.cell(row=r, column=1, value=t).font = F_BODY
    ws1.cell(row=r, column=1).alignment = A_LEFT
    ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws1.row_dimensions[r].height = 28

# 固定工作表顺序
ORDER = ["01-说明与查询", "02-基础参数表", "03-指标清单", "04-公式变量映射", "05-数据源清单", "06-枚举字典"]
wb._sheets = [wb[n] for n in ORDER]
wb.active = 0

wb.save(OUT)

# ============================================================
# 4. 为公式单元格注入缓存值（保证任何预览器都能正确显示）
# ============================================================
import zipfile, shutil, html


def inject_cached_values(path, sheet_title, values):
    """values: {cell_ref: value}，可为 str / int / float"""
    tmp = path + ".tmp"
    zin = zipfile.ZipFile(path, "r")
    names = zin.namelist()
    wbxml = zin.read("xl/workbook.xml").decode("utf-8")
    relxml = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    m = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(sheet_title), wbxml)
    if not m:
        m = re.search(r'<sheet[^>]*r:id="([^"]+)"[^>]*name="%s"' % re.escape(sheet_title), wbxml)
    rid = m.group(1)
    rel = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % re.escape(rid), relxml)
    if not rel:
        rel = re.search(r'Target="([^"]+)"[^>]*Id="%s"' % re.escape(rid), relxml)
    target = rel.group(1).lstrip("/")
    if not target.startswith("xl/"):
        target = "xl/" + target
    xml = zin.read(target).decode("utf-8")
    for ref, val in values.items():
        pat = re.compile(r'(<c r="%s"(?:[^>]*?)>)(<f>.*?</f>)(?:<v>.*?</v>)?(</c>)' % ref, re.S)

        def repl(mm):
            head, fx, tail = mm.group(1), mm.group(2), mm.group(3)
            if isinstance(val, str):
                head = re.sub(r'\s+t="[^"]*"', "", head)
                head = head[:-1] + ' t="str">'
                v = "<v>%s</v>" % html.escape(val, quote=False)
            else:
                head = re.sub(r'\s+t="[^"]*"', "", head)
                v = "<v>%s</v>" % val
            return head + fx + v + tail
        xml, n = pat.subn(repl, xml, count=1)
        if n == 0:
            raise RuntimeError("未找到公式单元格 %s" % ref)
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    for nm in names:
        data = zin.read(nm)
        if nm == target:
            data = xml.encode("utf-8")
        zout.writestr(nm, data)
    zout.close()
    zin.close()
    shutil.move(tmp, path)


uni_by_code = {row[0]: row for row in UNIFORM}
target_row = uni_by_code["EN-013"]
cached = {"D6": "匹配成功",
          "B7": 3 + [r[0] for r in UNIFORM].index("EN-013")}
for i, (label, col) in enumerate(QUERY_FIELDS):
    v = target_row[COLMAP[col] - 2]
    cached["B%d" % (start + i)] = v if isinstance(v, (int, float)) else str(v)
inject_cached_values(OUT, "01-说明与查询", cached)
print("cached values injected:", len(cached))

print("OK ->", OUT)
print("uniform rows:", len(UNIFORM), "| indicators:", len(INDICATORS), "| params:", len(PARAMS),
      "| var rows:", len(rows4), "| sources:", len(SOURCES), "| enums:", len(ENUMS), "| last row:", last2)
