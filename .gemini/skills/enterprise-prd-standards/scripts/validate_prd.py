# -*- coding: utf-8 -*-
"""
企业级 PRD 质量与工程完整性静态验证脚本 (validate_prd.py)
用于扫描 Markdown 或 Word (.docx) PRD 文档中的占位符、敏感定性词汇、八段式结构完整度与关键技术契约。

使用方法:
    py validate_prd.py <path_to_prd_file>
"""

import os
import sys
import re

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

FORBIDDEN_PLACEHOLDERS = [
    "待补充", "待确认", "TBD", "FIXME", "TODO", "xxx", "XXX", "暂定", "待完善", "待讨论"
]

SENSITIVE_QUALITATIVE_WORDS = [
    "表现优异", "落后单位", "运行欠佳", "电能品质优良", "达标奖励", "处于领跑标杆"
]

EIGHT_SECTIONS = [
    "【1】", "【2】", "【3】", "【4】", "【5】", "【6】", "【7】", "【8】"
]

def load_file_content(filepath):
    if not os.path.exists(filepath):
        print(f"[ERROR] 文件不存在: {filepath}")
        sys.exit(1)
        
    ext = os.path.splitext(filepath)[1].lower()
    if ext in [".md", ".txt", ".html"]:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif ext in [".docx"]:
        try:
            import docx
            doc = docx.Document(filepath)
            lines = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    for c in row.cells:
                        lines.append(c.text)
            return "\n".join(lines)
        except ImportError:
            print("[WARN] 未安装 python-docx，无法解析 docx 文件，请先安装 pip install python-docx")
            sys.exit(1)
    else:
        print(f"[ERROR] 不支持的文件扩展名: {ext}")
        sys.exit(1)

def validate_prd(filepath):
    print("=" * 70)
    print(f"  企业级 PRD 质量与工程完整性自动化静态审计")
    print(f"  目标文件: {filepath}")
    print("=" * 70)

    content = load_file_content(filepath)
    issues_count = 0

    # 1. 占位符扫描 (Zero Tolerance)
    print("\n[门禁 1] 未决占位符扫描 (DoR 一票否决项)...")
    placeholder_hits = {}
    for kw in FORBIDDEN_PLACEHOLDERS:
        cnt = content.count(kw)
        if cnt > 0:
            placeholder_hits[kw] = cnt
            issues_count += cnt

    if placeholder_hits:
        print("  ❌ [FAIL] 发现未决占位符残留：")
        for kw, cnt in placeholder_hits.items():
            print(f"     - '{kw}': 出现 {cnt} 次")
    else:
        print("  ✅ [PASS] 未决占位符完全归零 (0 残留)。")

    # 2. 主观定性词汇扫描 (客观中立原则)
    print("\n[门禁 2] 主观定性褒贬评价词汇扫描 (客观中立准则)...")
    qualitative_hits = {}
    for kw in SENSITIVE_QUALITATIVE_WORDS:
        cnt = content.count(kw)
        if cnt > 0:
            qualitative_hits[kw] = cnt

    if qualitative_hits:
        print("  ⚠️ [WARN] 发现潜在的主观定性词汇（需核验是否仅作为测试拦截靶点）：")
        for kw, cnt in qualitative_hits.items():
            print(f"     - '{kw}': 出现 {cnt} 次")
    else:
        print("  ✅ [PASS] 0 主观定性词汇，文本纯洁度符合客观中立原则。")

    # 3. 八段式结构完整度检查
    print("\n[门禁 3] 八段式模块结构完整度扫描...")
    sec_hits = {}
    for s in EIGHT_SECTIONS:
        cnt = content.count(s)
        sec_hits[s] = cnt

    missing_secs = [s for s, cnt in sec_hits.items() if cnt == 0]
    if missing_secs:
        print(f"  ⚠️ [INFO] 本文档未包含或部分缺失八段式标记: {', '.join(missing_secs)}")
    else:
        min_cnt = min(sec_hits.values())
        print(f"  ✅ [PASS] 八段式结构标记全量覆盖，共识别出至少 {min_cnt} 个标准模块。")

    # 4. 关键工程防御机制检查
    print("\n[门禁 4] 核心工程防御机制关键字检测...")
    defenses = {
        "除零安全兜底 (--)": "--" in content,
        "受控错误码 (E_...)": bool(re.search(r"E_[A-Z0-9_]+", content)),
        "BDD 验收用例 (Scenario/Given)": "Given" in content or "Scenario:" in content,
        "数学公式/模型": "$" in content or "公式" in content or "∑" in content or "Σ" in content,
    }
    for def_name, passed in defenses.items():
        status = "✅ [PASS]" if passed else "❌ [MISSING]"
        print(f"  {status} {def_name}")
        if not passed:
            issues_count += 1

    print("\n" + "=" * 70)
    if issues_count == 0:
        print("  🏆 最终审计裁决: PASS —— 本 PRD 具备高成熟度与研发施工准入条件！")
    else:
        print(f"  ⚠️ 最终审计裁决: NEED_REVISION —— 共发现 {issues_count} 项待优化或阻断项。")
    print("=" * 70)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: py validate_prd.py <path_to_prd_file>")
    else:
        validate_prd(sys.argv[1])
